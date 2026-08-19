import socket as sk
from email.mime import image
import sys
from unittest import result
import mediapipe as mp
import cv2
import os
import numpy as np
import math
from datetime import datetime

from openpyxl import Workbook, load_workbook
from calculation_amplitude import CalculationAmplitudeClass
from vector_drawer import VectorDrawer
#import json

#coff = np.array([2.145035e-03, -0.631834, 66.585278])
coff = np.array([1.405132e-03, -0.486281, 62.507305]) #calcule em calibrar.py
class HandDetection:
    # Inicializa a classe HandDetection
    def __init__(self, pairs, camera_index=0, min_detection_confidence=0.8, min_tracking_confidence=0.8): #TODO: MUDAR CONFIDENCE PARA .8 PARA UNITY
        self.frame_count = 0
        self.save_every = 30  # salva a cada 30 frames (~1 segundo se estiver a 30fps)
        self.excel_path = None  # caminho do arquivo será definido no primeiro save

        if pairs is not None:
            self.pairs = pairs
            self.angles_data = {pair: [] for pair in self.pairs}
        else:
            self.pairs = [("THUMB_TIP", "INDEX_FINGER_TIP")]

        # Dicionário que guarda a lista de ângulos de cada par
        self.angles_data = {pair: [] for pair in self.pairs}

        self.exibit = {
            "THUMB_TIP": "Polegar",
            "INDEX_FINGER_TIP": "Indicador",
            "MIDDLE_FINGER_TIP": "Medio",
            "RING_FINGER_TIP": "Anelar",
            "PINKY_TIP": "Mindinho"
        }

        #cofiguracao da comunicacao UDP para o unity
        self.sock = sk.socket(sk.AF_INET, sk.SOCK_DGRAM)
        self.serverAddrPlusPort = ("127.0.0.1", 5052) #usar do windows para wsl

        self.mp_selfie_segmentation = mp.solutions.selfie_segmentation.SelfieSegmentation(model_selection=1)

        # Detecta apenas uma mão.
        self.mp_drawing = mp.solutions.drawing_utils
        self.mp_hands = mp.solutions.hands.Hands(max_num_hands=1, min_detection_confidence=min_detection_confidence,
                                                 min_tracking_confidence=min_tracking_confidence, model_complexity=1)
        self.cap = cv2.VideoCapture(camera_index)
        self.calc_amplitude = CalculationAmplitudeClass()
        self.vector_drawer = VectorDrawer()

    def get_distance(self, hand_landmarks, height, width):
        """
                Calcula a distância real da mão até a câmera em cm usando regressão quadrática.
                Retorna: (distance_cm, lmList)
                """
        lmList = []
        for i in range(21):
            landmark = hand_landmarks.landmark[i]
            x = landmark.x * width
            y = height - (landmark.y * height)
            z = landmark.z * width
            lmList.extend([x, y, z])

        # Pontos: base do indicador (5) e base do mindinho (17)
        x1, y1 = lmList[5 * 3], lmList[5 * 3 + 1]
        x2, y2 = lmList[17 * 3], lmList[17 * 3 + 1]
        # Distância em pixels
        distance_pixels = int(math.hypot(x2 - x1, y2 - y1))
        # Aplicar regressão quadrática
        A, B, C = coff
        distance_cm = A * distance_pixels ** 2 + B * distance_pixels + C

        # Limitar a faixa física realista
        distance_cm = max(20.0, min(100.0, distance_cm))

        return distance_cm, lmList

    # Processa cada frame do programa enquanto a câmera estiver ativa
    def process_frame(self):
        ret, frame = self.cap.read()
        if not ret:
            return None, None

        image = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        image.flags.writeable = False
        results = self.mp_hands.process(image)
        image.flags.writeable = True
        image = cv2.cvtColor(image, cv2.COLOR_RGB2BGR)
        return image, results

    def background_color(self, image, results=None):
        # Converte para RGB, que é o formato que o MediaPipe espera
        image_rgb = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)

        # Segmenta a imagem com MediaPipe Selfie Segmentation
        segmentation_results = self.mp_selfie_segmentation.process(image_rgb)

        # A segmentação gera uma máscara com valores de confiança entre 0 e 1
        condition = segmentation_results.segmentation_mask > 0.8  # threshold

        # Cria fundo colorido
        bg_color = (255, 255, 255)  # branco
        bg_image = np.full(image.shape, bg_color, dtype=np.uint8)

        # Combina imagem original e fundo com base na máscara
        final_image = np.where(condition[..., None], image, bg_image)

        return final_image

    def get_hand_bbox(self, hand_landmarks, img_width, img_height):
        """Retorna bounding box da mão em pixels (x, y, w, h)"""
        x_coords = [lm.x * img_width for lm in hand_landmarks.landmark]
        y_coords = [lm.y * img_height for lm in hand_landmarks.landmark]
        x_min, x_max = int(min(x_coords)), int(max(x_coords))
        y_min, y_max = int(min(y_coords)), int(max(y_coords))
        return x_min, y_min, x_max - x_min, y_max - y_min
    # Desenha os pontos das mãos detectados
    def draw_landmarks(self, image, results):
        height, width = image.shape[:2]
        if not results.multi_hand_landmarks:
            return
        for hand_landmarks in results.multi_hand_landmarks:

            self.mp_drawing.draw_landmarks(image, hand_landmarks, mp.solutions.hands.HAND_CONNECTIONS)
            wrist = hand_landmarks.landmark[mp.solutions.hands.HandLandmark.WRIST]

            y_offset = 30
            for pair in self.pairs:
                finger1, finger2 = pair
                landmrk1 = hand_landmarks.landmark[getattr(mp.solutions.hands.HandLandmark, finger1)]
                landmrk2 = hand_landmarks.landmark[getattr(mp.solutions.hands.HandLandmark, finger2)]

                # Desenha os vetores formados pelos membros
                self.vector_drawer.draw_vector(image, wrist, landmrk1, color=(0, 255, 0))
                self.vector_drawer.draw_vector(image, wrist, landmrk2, color=(0, 255, 0))

                # Calcula o ângulo e a amplitude de movimento
                vector_1 = self.calc_amplitude.create_vector(wrist, landmrk1)
                vector_2 = self.calc_amplitude.create_vector(wrist, landmrk2)
                angle = self.calc_amplitude.calculate_amplitude(vector_1, vector_2)

                # Guarda o valor
                self.angles_data[pair].append(angle)

                # Exibe na tela resultados
                name1= self.exibit.get(finger1, finger1)
                name2 = self.exibit.get(finger2, finger2)
                cv2.putText(image, f'{name1}↔{name2}: {angle:.1f}', (10, y_offset), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (255, 255, 255), 2)
                y_offset += 30

    def d2_to_unity(self, image, results):
        height, width = image.shape[:2]
        if results.multi_hand_landmarks:
            for hand_landmarks in results.multi_hand_landmarks:
                # === 1. CALCULAR DISTÂNCIA EM CM + LANDMARKS ===
                distance_cm, lmList = self.get_distance(hand_landmarks, height, width)
                Z_meters = -distance_cm / 100.0  # em metros, negativo = frente

                # === 2. MONTAR MENSAGEM: 63 valores + Z em metros ===
                message_parts = [f"{v:.3f}" for v in lmList]  # 63 valores
                message_parts.append(f"{Z_meters:.3f}")  # último: Z em metros
                message = ','.join(message_parts)

                # === 3. ENVIAR PARA UNITY ===
                self.sock.sendto(message.encode('utf-8'), self.serverAddrPlusPort)
                print(message[:200] + "..." + message[-50:])  # opcional: ver começo e fim

    def d3_to_unity(self, image, results):
        if not results.multi_hand_landmarks or not results.multi_hand_world_landmarks:
            return

        height, width = image.shape[:2]

        # Usamos os dois tipos de landmarks juntos
        for img_hand, world_hand in zip(results.multi_hand_landmarks, results.multi_hand_world_landmarks):
            wrist_img = img_hand.landmark[0]

            X_offset = (wrist_img.x - 0.5) * 2.0  # centraliza (-1 a 1)
            Y_offset = (wrist_img.y - 0.5) * -2.0  # inverte eixo Y

            # Calcula distância usando o landmark da IMAGEM (correto)
            distance_cm, _ = self.get_distance(img_hand, height, width)
            Z_offset = -distance_cm / 100.0

            # Monta os 63 valores usando o WORLD landmark
            lm_3d_list = []
            for lm in world_hand.landmark:
                lm_3d_list.extend([lm.x, lm.y, lm.z])

            # Adiciona o Z calibrado no final
            message_parts = [f"{v:.3f}" for v in lm_3d_list]
            message_parts.append(f"{X_offset:.3f}")
            message_parts.append(f"{Y_offset:.3f}")
            message_parts.append(f"{Z_offset:.3f}")

            message = ','.join(message_parts)
            self.sock.sendto(message.encode('utf-8'), self.serverAddrPlusPort)
            print(message[:200] + "..." + message[-50:])   # opcional

    def save_to_excel(self):
        # Descobre o diretório onde está o .exe ou .py
        if getattr(sys, 'frozen', False):
            BASE_DIR = os.path.dirname(sys.executable)  # Se for executável
        else:
            BASE_DIR = os.path.dirname(os.path.abspath(__file__))  # Se for script

        # Pasta de saída
        PASTA_DATAS = os.path.join(BASE_DIR, "datas")
        os.makedirs(PASTA_DATAS, exist_ok=True)

        # Criar o caminho do arquivo com data
        if self.excel_path is None:
            timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            self.excel_path = os.path.join(PASTA_DATAS, f"angles_{timestamp}.xlsx")

            wb = Workbook()
            wb.remove(wb.active) #remove aba padrao

            for pair in self.pairs:
                name1 = self.exibit.get(pair[0], pair[0])
                name2 = self.exibit.get(pair[1], pair[1])
                sheet_name = f"{name1}-{name2}"[:31]
                ws = wb.create_sheet(title=sheet_name)
                ws['A1'] = f"Angulo entre {name1} e {name2}"
                ws['A2'] = "Angulo (graus)"
            wb.save(self.excel_path)

        wb= load_workbook(self.excel_path)

        for pair in self.pairs:
            name1 = self.exibit.get(pair[0], pair[0])
            name2 = self.exibit.get(pair[1], pair[1])
            sheet_name = f"{name1}-{name2}"[:31]
            ws = wb[sheet_name]

            # Apaga os dados antigos (mantém só o cabeçalho)
            if ws.max_row > 2:
                ws.delete_rows(3, ws.max_row - 2)

            # Escreve todos os ângulos atuais
            for i, angle in enumerate(self.angles_data[pair], start=3):
                ws.cell(row=i, column=1, value=round(angle, 2))

        wb.save(self.excel_path)
        print(f"Excel atualizado: {self.excel_path}")


    # Executa a detecção da mão através da câmera
    def run(self):
        contador = 1

        while self.cap.isOpened():
            image, results = self.process_frame()
            if image is None:
                break

            #converte landmarks em uma lista para serem passadas ao unity dentro do drawlandmarks
            # Desenha os vetores e ângulo sobre a imagem final
            self.draw_landmarks(image, results)

            # mover logica errada em draw_landmarks para ca
            #self.d2_to_unity(image, results)
            self.d3_to_unity(image, results)

            # Salva periodicamente
            self.frame_count += 1
            if self.frame_count % self.save_every == 0:
                self.save_to_excel()

            cv2.imshow('Hand Detection', image)
            if cv2.waitKey(10) & 0xFF == 27: #ESC
                contador += 1  # incrementa para o próximo arquivo
                break

        self.cap.release()
        cv2.destroyAllWindows()

        # Salva tudo no Excel ao sair
        self.save_to_excel()